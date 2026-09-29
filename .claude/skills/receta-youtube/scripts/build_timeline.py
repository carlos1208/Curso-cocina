#!/usr/bin/env python3
"""guion.json + hoja_tomas.xlsx → timeline.json + subtitulos.srt (the long 16:9 cut).

The sheet owns the timing (sheet "Tomas": one row per shot, cut on the narration's pauses); guion.json owns the design
(what each shot shows, callouts, pins, subtitles, hook, end screen). Anywhere in guion.json a time can be written as
"@S07" (start of shot S07) or "@S07$" (its end), so moving a cut in the sheet moves everything tied to it.
Checks before you render: duration within the target, every shot has a source, clip ranges fit their clips.
    python3 build_timeline.py [guion.json]
"""
import json
import os
import re
import sys

import openpyxl

G = json.load(open(sys.argv[1] if len(sys.argv) > 1 else 'guion.json', encoding='utf-8'))
LIMITS = G.get('duration_limits', [150, 240])           # 2.5–4 min

# ── timing from the sheet (columns found by header, so extra columns don't matter) ──
wb = openpyxl.load_workbook(G.get('sheet', 'hoja_tomas.xlsx'), data_only=True)
ws = wb['Tomas'] if 'Tomas' in wb.sheetnames else wb.worksheets[0]
rows = list(ws.iter_rows(values_only=True))
head = [str(h or '').strip().lower() for h in rows[0]]
col = lambda *names: next(i for i, h in enumerate(head) if any(n in h for n in names))
ci, c0, c1 = 0, col('inicio (s)', 'inicio'), col('fin (s)', 'fin')
T = {}
for r in rows[1:]:
    if r[ci] and re.match(r'^S\d+', str(r[ci])):
        T[str(r[ci])] = (float(r[c0]), float(r[c1]))
ORDER = sorted(T, key=lambda k: T[k][0])
DUR = round(max(t1 for _, t1 in T.values()), 2)
warn = []
if not LIMITS[0] <= DUR <= LIMITS[1]:
    warn.append(f'duration {DUR:.1f} s is outside {LIMITS[0]}–{LIMITS[1]} s')
for a, b in zip(ORDER, ORDER[1:]):
    if abs(T[a][1] - T[b][0]) > 0.05:
        warn.append(f'gap/overlap between {a} and {b}: {T[a][1]} → {T[b][0]}')


def resolve(x):
    """'@S07' → start of S07, '@S07$' → end of S07 (recursively)."""
    if isinstance(x, str):
        m = re.fullmatch(r'@(S\d+)(\$?)', x)
        if m:
            if m.group(1) not in T:
                sys.exit(f'unknown shot {m.group(1)} in {x}')
            return T[m.group(1)][1 if m.group(2) else 0]
        return x
    if isinstance(x, list):
        return [resolve(v) for v in x]
    if isinstance(x, dict):
        return {k: resolve(v) for k, v in x.items()}
    return x


G = resolve(G)

# ── media: clips are read from frames/<key>/ (prep_media.py), stills straight from their path ──
index = json.load(open('frames/index.json')) if os.path.exists('frames/index.json') else {}
MEDIA = {}
for k, m in G['media'].items():
    MEDIA[k] = {'img': m['img']} if 'img' in m else {'clip': k, 'audio': m.get('audio', True)}

# ── chapters: hook at 0, cooking chapters alternate the card's side, end screen last ──
CH = [dict(id='gancho', t=0.0, n=0, title=G['title'], color=G['chapters'][0]['color'], side='hook')]
for i, c in enumerate(G['chapters']):
    CH.append(dict(id=c.get('id', f'cap{i + 1}'), t=T[c['start']][0] if c['start'] in T else c['start'], n=i + 1, title=c['title'],
                   h1=c.get('h1', c['title']), h2=c.get('h2', ''), color=c['color'], side='R' if i % 2 == 0 else 'L', mark=c.get('mark', '')))
if G.get('end'):
    CH.append(dict(id='cierre', t=T[G['end']['start']][0], n=len(CH), title='Cierre', color=G['end'].get('color', '#F2B544'), side='end'))
starts = {c['t'] for c in CH[1:]}

# ── shots ──
SHOTS = []
prev = None
for sid in ORDER:
    t0, t1 = T[sid]
    spec = G['shots'].get(sid)
    if not spec:
        sys.exit(f'{sid} has no entry in guion.json "shots"')
    src = spec['src']
    if src not in MEDIA:
        sys.exit(f'{sid}: source {src} is not in "media"')
    a = spec.get('a', 0.0)
    b = spec.get('b', a + (t1 - t0))
    if 'clip' in MEDIA[src] and src in index:
        length = index[src]['frames'] / index[src]['fps']
        if b > length + 0.05:
            warn.append(f'{sid}: {src} is {length:.2f} s long but the shot asks for {a:.2f}–{b:.2f} s')
    if t0 < CH[1]['t']:
        tr = 'none'
    elif CH[-1]['side'] == 'end' and t0 >= CH[-1]['t']:
        tr = 'end'
    elif t0 in starts:
        tr = 'chapter'
    elif prev and prev['src'] == src and 'img' in MEDIA[src]:
        tr = 'zoom'
    else:
        tr = 'push'
    s = dict(id=sid, t0=t0, t1=t1, src=src, a=a, b=b, z=spec.get('z', [1.03, 1.08]),
             f=spec.get('f', [[.5, .5], [.5, .5]]), tr=spec.get('tr', tr))
    if spec.get('shimmer'):
        s['shimmer'] = spec['shimmer']
    SHOTS.append(s); prev = s

# ── callouts ("until": shot id → ends with that shot; default: the shot where it starts) ──
def shot_at(t):
    return next((s for s in SHOTS if s['t0'] - 1e-6 <= t < s['t1']), SHOTS[-1])


CO = []
for o in G.get('callouts', []):
    o = dict(o)
    if 'until' in o:
        o['t1'] = T[o.pop('until')][1]
    o.setdefault('t1', shot_at(o['t'])['t1'])
    CO.append(o)
PINS = G.get('pins', [])


# ── subtitles: the corrected lines, joined into readable cues ──
def cues(subs, max_chars=92, gap=0.35):
    out = []
    plain = lambda x: x.replace('**', '')
    for t0, t1, s in subs:
        if out:
            p = out[-1]
            if t0 - p[1] < gap and not plain(p[2]).rstrip().endswith(('.', '?', '!', '…')) and len(plain(p[2] + ' ' + s)) <= max_chars:
                p[1], p[2] = t1, p[2] + ' ' + s
                continue
        out.append([t0, t1, s])
    for i in range(len(out) - 1):
        if out[i + 1][0] - out[i][1] < 0.6:
            out[i][1] = out[i + 1][0] - 0.02
    return out


CUES = cues(G['subtitles'])


def srt_time(t):
    ms = int(round(t * 1000)); h, ms = divmod(ms, 3600000); m, ms = divmod(ms, 60000); s, ms = divmod(ms, 1000)
    return f'{h:02d}:{m:02d}:{s:02d},{ms:03d}'


with open('subtitulos.srt', 'w', encoding='utf-8') as fh:
    for i, (t0, t1, s) in enumerate(CUES, 1):
        fh.write(f'{i}\n{srt_time(t0)} --> {srt_time(t1)}\n{s.replace("**", "")}\n\n')

TLJ = dict(duration=DUR, fps=30, title=G['title'], tag=G.get('tag', ''), music=G.get('music', 'latino'),
           colors=G.get('palette', {}), chapters=CH, media=MEDIA, shots=SHOTS, callouts=CO, pins=PINS, cues=CUES,
           hook=G.get('hook', {}), end=G.get('end', {}), thumb=G.get('thumb', {}))
if G.get('palette'):
    TLJ['palette'] = G['palette']
json.dump(TLJ, open('timeline.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
print(f'timeline.json: {DUR:.2f} s · {len(SHOTS)} shots · {len(CH)} chapters · {len(CO)} callouts · {len(CUES)} subtitle cues')
for w in warn:
    print('  ! ' + w)
